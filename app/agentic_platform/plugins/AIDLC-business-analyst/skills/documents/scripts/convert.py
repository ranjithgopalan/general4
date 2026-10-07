#!/usr/bin/env python3
"""
Document Conversion Utility

Converts between Markdown and various document formats:
- DOCX (Word)
- XLSX (Excel)
- PDF
- PPTX (PowerPoint)
- MSG (Outlook email)
"""

import argparse
import re
import subprocess
import sys
from pathlib import Path
from typing import Optional, List, Dict, Tuple

# Dependency mapping: command -> required packages
# Note: md2docx delegates to GATHER-help:md-to-word skill (no dependencies here)
COMMAND_DEPENDENCIES: Dict[str, List[tuple]] = {
    'docx2md': [('python-docx', 'docx')],
    'table2xlsx': [('openpyxl', 'openpyxl')],
    'xlsx2table': [('openpyxl', 'openpyxl')],
    'xlsx2csv': [('openpyxl', 'openpyxl')],
    'csv2xlsx': [('openpyxl', 'openpyxl')],
    'md2pdf': [('weasyprint', 'weasyprint'), ('markdown', 'markdown')],
    'pdf2md': [('pdfplumber', 'pdfplumber')],
    'md2pptx': [('python-pptx', 'pptx')],
    'pptx2md': [('python-pptx', 'pptx')],
    'msg2md': [('extract-msg', 'extract_msg'), ('beautifulsoup4', 'bs4')],
    'msg2eml': [('extract-msg', 'extract_msg')],
    'msg-extract': [('extract-msg', 'extract_msg')],
}


def ensure_dependencies(command: str) -> None:
    """Check and install required dependencies for a command."""
    if command not in COMMAND_DEPENDENCIES:
        return

    for package_name, import_name in COMMAND_DEPENDENCIES[command]:
        try:
            __import__(import_name)
        except ImportError:
            print(f"Installing {package_name}...")
            subprocess.check_call(
                [sys.executable, '-m', 'pip', 'install', '-q', package_name],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL
            )
            print(f"Installed {package_name}")


def parse_a1_notation(notation: str) -> Tuple[Optional[str], int, int, Optional[int], Optional[int]]:
    """
    Parse A1 notation like 'Sheet1!A1:D10' or 'B2:F20' or 'A1'.
    Returns: (sheet_name, start_row, start_col, end_row, end_col)
    Row/col are 1-indexed. end_row/end_col are None if single cell.
    """
    sheet_name = None
    if '!' in notation:
        sheet_name, notation = notation.split('!', 1)

    def col_to_num(col: str) -> int:
        """Convert column letter(s) to number (A=1, Z=26, AA=27)."""
        result = 0
        for char in col.upper():
            result = result * 26 + (ord(char) - ord('A') + 1)
        return result

    def parse_cell(cell: str) -> Tuple[int, int]:
        """Parse cell reference like 'A1' to (row, col)."""
        match = re.match(r'^([A-Za-z]+)(\d+)$', cell)
        if not match:
            raise ValueError(f"Invalid cell reference: {cell}")
        col = col_to_num(match.group(1))
        row = int(match.group(2))
        return row, col

    if ':' in notation:
        start, end = notation.split(':')
        start_row, start_col = parse_cell(start)
        end_row, end_col = parse_cell(end)
        return sheet_name, start_row, start_col, end_row, end_col
    else:
        row, col = parse_cell(notation)
        return sheet_name, row, col, None, None


def num_to_col(n: int) -> str:
    """Convert column number to letter(s) (1=A, 26=Z, 27=AA)."""
    result = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        result = chr(65 + remainder) + result
    return result


class DocumentConverter:
    """Handles conversions between Markdown and various document formats."""

    def md_to_docx(self, input_path: str, output_path: str) -> None:
        """
        Convert Markdown file to DOCX using GATHER-help:md-to-word skill.

        This delegates to the md2word CLI tool from GATHER-help plugin for improved
        markdown to Word conversion with proper TOC support, table handling, and formatting.
        """
        import os

        # Validate input file
        if not Path(input_path).exists():
            print(f"Error: Input file not found: {input_path}")
            sys.exit(1)

        # Determine md2word CLI path (cross-platform)
        home = os.path.expanduser('~')
        if sys.platform == 'win32':
            md2word_path = os.path.join(home, '.claude', 'venvs', 'gather', 'Scripts', 'md2word')
        else:
            md2word_path = os.path.join(home, '.claude', 'venvs', 'gather', 'bin', 'md2word')

        # Check if md2word is available
        if not os.path.exists(md2word_path):
            print(f"Error: md2word CLI not found at {md2word_path}")
            print("Please ensure GATHER-help:md-to-word skill is installed.")
            sys.exit(1)

        # Call md2word CLI
        try:
            result = subprocess.run(
                [md2word_path, str(input_path), str(output_path)],
                capture_output=True,
                text=True,
                check=True
            )
            print(result.stdout)
        except subprocess.CalledProcessError as e:
            print(f"Error during conversion: {e}")
            if e.stderr:
                print(e.stderr)
            sys.exit(1)
        except FileNotFoundError:
            print(f"Error: Could not execute md2word at {md2word_path}")
            print("Please ensure GATHER-help:md-to-word skill is installed.")
            sys.exit(1)

    def docx_to_md(self, input_path: str, output_path: str) -> str:
        """Convert DOCX file to Markdown."""
        from docx import Document

        doc = Document(input_path)
        md_lines = []

        for element in doc.element.body:
            # Handle paragraphs
            if element.tag.endswith('p'):
                for para in doc.paragraphs:
                    if para._element == element:
                        style_name = para.style.name if para.style else ''
                        text = para.text.strip()

                        if not text:
                            md_lines.append('')
                            continue

                        # Handle headings
                        if 'Heading' in style_name:
                            try:
                                level = int(style_name.replace('Heading ', ''))
                                md_lines.append(f"{'#' * level} {text}")
                            except ValueError:
                                md_lines.append(f"# {text}")
                        # Handle lists
                        elif 'List Bullet' in style_name:
                            md_lines.append(f"- {text}")
                        elif 'List Number' in style_name:
                            md_lines.append(f"1. {text}")
                        else:
                            # Regular paragraph with formatting
                            formatted = self._extract_formatting(para)
                            md_lines.append(formatted)
                        break

            # Handle tables
            elif element.tag.endswith('tbl'):
                for table in doc.tables:
                    if table._element == element:
                        table_md = self._table_to_markdown(table)
                        md_lines.append(table_md)
                        break

        md_content = '\n'.join(md_lines)

        if output_path:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(md_content)
            print(f"Created: {output_path}")

        return md_content

    def table_to_xlsx(self, input_path: str, output_path: str) -> None:
        """Extract Markdown tables to XLSX."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter

        # Read markdown
        with open(input_path, 'r', encoding='utf-8') as f:
            md_content = f.read()

        # Find all tables
        tables = self._extract_markdown_tables(md_content)

        if not tables:
            print("No tables found in markdown file")
            return

        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        for idx, table_data in enumerate(tables):
            sheet_name = f"Table{idx + 1}"
            ws = wb.create_sheet(title=sheet_name)

            for row_idx, row in enumerate(table_data, start=1):
                for col_idx, cell_value in enumerate(row, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)

                    # Bold header row
                    if row_idx == 1:
                        cell.font = Font(bold=True)
                    cell.alignment = Alignment(wrap_text=True)

            # Auto-adjust column widths
            for col_idx in range(1, len(table_data[0]) + 1 if table_data else 1):
                col_letter = get_column_letter(col_idx)
                max_length = 0
                for row in table_data:
                    if col_idx <= len(row):
                        max_length = max(max_length, len(str(row[col_idx - 1])))
                ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

        wb.save(output_path)
        print(f"Created: {output_path} ({len(tables)} table(s))")

    def xlsx_to_markdown_table(
        self,
        input_path: str,
        output_path: Optional[str] = None,
        sheets: Optional[List[str]] = None,
        range_notation: Optional[str] = None,
        no_header: bool = False
    ) -> str:
        """
        Convert XLSX to Markdown tables.

        Args:
            input_path: Path to Excel file
            output_path: Output markdown file (optional)
            sheets: List of sheet names to include (None = all)
            range_notation: A1 notation like 'A1:D10' or 'Sheet1!B2:F20'
            no_header: If True, don't treat first row as header
        """
        from openpyxl import load_workbook

        wb = load_workbook(input_path)
        md_tables = []

        # Parse range if provided
        range_sheet = None
        start_row = start_col = 1
        end_row = end_col = None
        if range_notation:
            range_sheet, start_row, start_col, end_row, end_col = parse_a1_notation(range_notation)

        # Determine which sheets to process
        if range_sheet:
            sheet_list = [range_sheet]
        elif sheets:
            sheet_list = [s for s in sheets if s in wb.sheetnames]
        else:
            sheet_list = wb.sheetnames

        for sheet_name in sheet_list:
            if sheet_name not in wb.sheetnames:
                print(f"Warning: Sheet '{sheet_name}' not found")
                continue

            ws = wb[sheet_name]

            # Get cell range
            if end_row and end_col:
                rows = []
                for r in range(start_row, end_row + 1):
                    row_data = []
                    for c in range(start_col, end_col + 1):
                        cell_value = ws.cell(row=r, column=c).value
                        row_data.append(cell_value)
                    rows.append(tuple(row_data))
            else:
                rows = list(ws.iter_rows(values_only=True))
                # Apply start position offset if specified
                if start_row > 1 or start_col > 1:
                    rows = rows[start_row - 1:]
                    rows = [r[start_col - 1:] for r in rows]

            if not rows:
                continue

            # Filter out completely empty rows
            rows = [r for r in rows if any(c is not None for c in r)]
            if not rows:
                continue

            # Determine max columns
            max_cols = max(len(r) for r in rows)

            # Build table
            table_lines = []
            if len(sheet_list) > 1:
                table_lines.append(f"### {sheet_name}\n")

            if no_header:
                # Generate column headers (A, B, C, ...)
                header_cells = [num_to_col(i) for i in range(start_col, start_col + max_cols)]
                table_lines.append('| ' + ' | '.join(header_cells) + ' |')
                table_lines.append('|' + '|'.join(['---'] * max_cols) + '|')
                data_rows = rows
            else:
                # First row is header
                header = rows[0]
                header_cells = [str(c) if c is not None else '' for c in header]
                header_cells.extend([''] * (max_cols - len(header_cells)))
                table_lines.append('| ' + ' | '.join(header_cells) + ' |')
                table_lines.append('|' + '|'.join(['---'] * max_cols) + '|')
                data_rows = rows[1:]

            # Data rows
            for row in data_rows:
                cells = [str(c) if c is not None else '' for c in row]
                cells.extend([''] * (max_cols - len(cells)))
                # Escape pipe characters
                cells = [c.replace('|', '\\|') for c in cells]
                table_lines.append('| ' + ' | '.join(cells) + ' |')

            md_tables.append('\n'.join(table_lines))

        md_content = '\n\n'.join(md_tables)

        if output_path:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(md_content)
            print(f"Created: {output_path}")

        return md_content

    def list_sheets(self, input_path: str) -> List[str]:
        """List all sheet names in an Excel file."""
        from openpyxl import load_workbook
        wb = load_workbook(input_path, read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return sheets

    def get_sheet_info(self, input_path: str, sheet_name: Optional[str] = None) -> Dict:
        """Get information about sheet(s) including dimensions and data range."""
        from openpyxl import load_workbook

        wb = load_workbook(input_path, read_only=True)
        info = {}

        sheets_to_check = [sheet_name] if sheet_name else wb.sheetnames

        for name in sheets_to_check:
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            info[name] = {
                'dimensions': ws.dimensions,
                'max_row': ws.max_row,
                'max_column': ws.max_column,
                'max_column_letter': num_to_col(ws.max_column) if ws.max_column else 'A'
            }

        wb.close()
        return info

    def xlsx_to_csv(
        self,
        input_path: str,
        output_path: Optional[str] = None,
        sheet: Optional[str] = None,
        range_notation: Optional[str] = None,
        delimiter: str = ','
    ) -> str:
        """
        Convert Excel sheet to CSV.

        Args:
            input_path: Path to Excel file
            output_path: Output CSV file (optional, prints to stdout if None)
            sheet: Sheet name to convert (None = first sheet)
            range_notation: A1 notation like 'A1:D10' or 'Sheet1!B2:F20'
            delimiter: CSV delimiter (default: comma)
        """
        import csv
        from io import StringIO
        from openpyxl import load_workbook

        wb = load_workbook(input_path)

        # Parse range if provided
        range_sheet = None
        start_row = start_col = 1
        end_row = end_col = None
        if range_notation:
            range_sheet, start_row, start_col, end_row, end_col = parse_a1_notation(range_notation)

        # Determine which sheet to use
        if range_sheet:
            sheet_name = range_sheet
        elif sheet:
            sheet_name = sheet
        else:
            sheet_name = wb.sheetnames[0]

        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")

        ws = wb[sheet_name]

        # Get data
        rows = []
        if end_row and end_col:
            for r in range(start_row, end_row + 1):
                row_data = []
                for c in range(start_col, end_col + 1):
                    cell_value = ws.cell(row=r, column=c).value
                    row_data.append(cell_value if cell_value is not None else '')
                rows.append(row_data)
        else:
            for row in ws.iter_rows(min_row=start_row, min_col=start_col, values_only=True):
                rows.append([c if c is not None else '' for c in row])

        wb.close()

        # Write CSV
        output = StringIO()
        writer = csv.writer(output, delimiter=delimiter)
        writer.writerows(rows)
        csv_content = output.getvalue()

        if output_path:
            with open(output_path, 'w', encoding='utf-8', newline='') as f:
                f.write(csv_content)
            print(f"Created: {output_path}")
        else:
            print(csv_content)

        return csv_content

    def csv_to_xlsx(
        self,
        input_path: str,
        output_path: str,
        sheet_name: str = 'Sheet1',
        delimiter: str = ',',
        has_header: bool = True
    ) -> None:
        """
        Convert CSV to Excel.

        Args:
            input_path: Path to CSV file
            output_path: Output Excel file
            sheet_name: Name for the worksheet
            delimiter: CSV delimiter (default: comma)
            has_header: If True, bold the first row
        """
        import csv
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Read CSV
        with open(input_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f, delimiter=delimiter)
            for row_idx, row in enumerate(reader, start=1):
                for col_idx, value in enumerate(row, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    # Bold header row
                    if has_header and row_idx == 1:
                        cell.font = Font(bold=True)

        # Auto-adjust column widths
        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_length = 0
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
                for cell in row:
                    if cell:
                        max_length = max(max_length, len(str(cell)))
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

        wb.save(output_path)
        print(f"Created: {output_path}")

    # ==================== PDF Conversions ====================

    def md_to_pdf(self, input_path: str, output_path: str) -> None:
        """Convert Markdown file to PDF."""
        import markdown
        from weasyprint import HTML, CSS

        # Read markdown content
        if Path(input_path).exists():
            with open(input_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
        else:
            md_content = input_path  # Treat as raw markdown string

        # Convert markdown to HTML
        html_content = markdown.markdown(
            md_content,
            extensions=['tables', 'fenced_code', 'codehilite', 'toc']
        )

        # Wrap in HTML document with styling
        full_html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                       line-height: 1.6; max-width: 800px; margin: 40px auto; padding: 0 20px; }}
                h1, h2, h3, h4, h5, h6 {{ margin-top: 1.5em; margin-bottom: 0.5em; }}
                h1 {{ font-size: 2em; border-bottom: 2px solid #333; padding-bottom: 0.3em; }}
                h2 {{ font-size: 1.5em; border-bottom: 1px solid #ccc; padding-bottom: 0.2em; }}
                code {{ background: #f4f4f4; padding: 2px 6px; border-radius: 3px; font-size: 0.9em; }}
                pre {{ background: #f4f4f4; padding: 16px; border-radius: 6px; overflow-x: auto; }}
                pre code {{ background: none; padding: 0; }}
                table {{ border-collapse: collapse; width: 100%; margin: 1em 0; }}
                th, td {{ border: 1px solid #ddd; padding: 8px 12px; text-align: left; }}
                th {{ background: #f4f4f4; font-weight: bold; }}
                blockquote {{ border-left: 4px solid #ddd; margin: 1em 0; padding-left: 1em; color: #666; }}
                ul, ol {{ padding-left: 2em; }}
            </style>
        </head>
        <body>
            {html_content}
        </body>
        </html>
        """

        HTML(string=full_html).write_pdf(output_path)
        print(f"Created: {output_path}")

    def pdf_to_md(self, input_path: str, output_path: str) -> str:
        """Extract text from PDF and convert to Markdown."""
        import pdfplumber

        md_lines = []

        with pdfplumber.open(input_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                # Extract text
                text = page.extract_text()
                if text:
                    md_lines.append(text)

                # Extract tables
                tables = page.extract_tables()
                for table in tables:
                    if table and len(table) > 0:
                        md_lines.append('')
                        # Header row
                        header = [str(c) if c else '' for c in table[0]]
                        md_lines.append('| ' + ' | '.join(header) + ' |')
                        md_lines.append('|' + '|'.join(['---'] * len(header)) + '|')
                        # Data rows
                        for row in table[1:]:
                            cells = [str(c).replace('|', '\\|') if c else '' for c in row]
                            md_lines.append('| ' + ' | '.join(cells) + ' |')
                        md_lines.append('')

                if page_num < len(pdf.pages):
                    md_lines.append('\n---\n')  # Page break

        md_content = '\n'.join(md_lines)

        if output_path:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(md_content)
            print(f"Created: {output_path}")

        return md_content

    # ==================== PowerPoint Conversions ====================

    def md_to_pptx(self, input_path: str, output_path: str) -> None:
        """Convert Markdown to PowerPoint presentation."""
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN

        # Read markdown content
        if Path(input_path).exists():
            with open(input_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
        else:
            md_content = input_path

        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 aspect ratio
        prs.slide_height = Inches(7.5)

        lines = md_content.split('\n')
        current_slide = None
        current_content = []

        def add_content_to_slide(slide, content_lines):
            """Add content lines to slide body."""
            if not content_lines or not slide:
                return

            # Find or create text box for content
            left = Inches(0.5)
            top = Inches(1.5)
            width = Inches(12.333)
            height = Inches(5.5)

            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.word_wrap = True

            first_para = True
            for line in content_lines:
                line = line.strip()
                if not line:
                    continue

                if first_para:
                    p = tf.paragraphs[0]
                    first_para = False
                else:
                    p = tf.add_paragraph()

                # Handle bullet points
                if line.startswith('- ') or line.startswith('* '):
                    p.text = line[2:]
                    p.level = 0
                elif line.startswith('  - ') or line.startswith('  * '):
                    p.text = line[4:]
                    p.level = 1
                elif line.startswith('    - ') or line.startswith('    * '):
                    p.text = line[6:]
                    p.level = 2
                else:
                    p.text = self._strip_formatting(line)

                p.font.size = Pt(18)

        for line in lines:
            # H1 = Title slide
            if line.startswith('# ') and not line.startswith('##'):
                # Save previous slide content
                if current_slide and current_content:
                    add_content_to_slide(current_slide, current_content)
                    current_content = []

                title_text = line[2:].strip()
                slide_layout = prs.slide_layouts[6]  # Blank
                current_slide = prs.slides.add_slide(slide_layout)

                # Add centered title
                left = Inches(0.5)
                top = Inches(3)
                width = Inches(12.333)
                height = Inches(1.5)
                txBox = current_slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                p = tf.paragraphs[0]
                p.text = self._strip_formatting(title_text)
                p.font.size = Pt(44)
                p.font.bold = True
                p.alignment = PP_ALIGN.CENTER

            # H2 = Section/content slide title
            elif line.startswith('## ') and not line.startswith('###'):
                # Save previous slide content
                if current_slide and current_content:
                    add_content_to_slide(current_slide, current_content)
                    current_content = []

                title_text = line[3:].strip()
                slide_layout = prs.slide_layouts[6]  # Blank
                current_slide = prs.slides.add_slide(slide_layout)

                # Add title at top
                left = Inches(0.5)
                top = Inches(0.5)
                width = Inches(12.333)
                height = Inches(1)
                txBox = current_slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                p = tf.paragraphs[0]
                p.text = self._strip_formatting(title_text)
                p.font.size = Pt(32)
                p.font.bold = True

            # H3+ = Sub-heading within slide
            elif line.startswith('###'):
                match = re.match(r'^(#{3,})\s+(.+)$', line)
                if match:
                    current_content.append(f"**{match.group(2)}**")

            # Regular content
            elif line.strip():
                current_content.append(line)

        # Save last slide content
        if current_slide and current_content:
            add_content_to_slide(current_slide, current_content)

        prs.save(output_path)
        print(f"Created: {output_path}")

    def pptx_to_md(self, input_path: str, output_path: str) -> str:
        """Extract content from PowerPoint to Markdown."""
        from pptx import Presentation

        prs = Presentation(input_path)
        md_lines = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_title = None
            slide_content = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if not text:
                            continue

                        # First text is usually the title
                        if slide_title is None and para.font.size and para.font.size.pt >= 24:
                            slide_title = text
                        elif para.level > 0:
                            indent = '  ' * para.level
                            slide_content.append(f"{indent}- {text}")
                        else:
                            slide_content.append(text)

                # Handle tables
                if shape.has_table:
                    table = shape.table
                    rows = []
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        rows.append(cells)

                    if rows:
                        max_cols = max(len(r) for r in rows)
                        slide_content.append('')
                        # Header
                        header = rows[0] + [''] * (max_cols - len(rows[0]))
                        slide_content.append('| ' + ' | '.join(header) + ' |')
                        slide_content.append('|' + '|'.join(['---'] * max_cols) + '|')
                        # Data
                        for row in rows[1:]:
                            cells = row + [''] * (max_cols - len(row))
                            slide_content.append('| ' + ' | '.join(cells) + ' |')
                        slide_content.append('')

            # Write slide to markdown
            if slide_title:
                md_lines.append(f"## {slide_title}")
            else:
                md_lines.append(f"## Slide {slide_num}")

            md_lines.append('')
            md_lines.extend(slide_content)
            md_lines.append('')

            # Check for speaker notes
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    md_lines.append(f"> **Notes:** {notes_text}")
                    md_lines.append('')

        md_content = '\n'.join(md_lines)

        if output_path:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(md_content)
            print(f"Created: {output_path}")

        return md_content

    # ==================== MSG (Outlook) Conversions ====================

    def msg_to_md(self, input_path: str, output_path: str) -> str:
        """Extract Outlook MSG file to Markdown."""
        import extract_msg
        from bs4 import BeautifulSoup

        msg = extract_msg.Message(input_path)

        md_lines = []

        # Title (Subject)
        md_lines.append(f"# {msg.subject or '(No Subject)'}")
        md_lines.append('')

        # Email headers table
        md_lines.append('| Field | Value |')
        md_lines.append('|-------|-------|')
        md_lines.append(f"| From | {msg.sender or 'Unknown'} |")
        md_lines.append(f"| To | {msg.to or ''} |")
        if msg.cc:
            md_lines.append(f"| CC | {msg.cc} |")
        if msg.bcc:
            md_lines.append(f"| BCC | {msg.bcc} |")
        md_lines.append(f"| Date | {msg.date or ''} |")
        md_lines.append('')

        # Body
        md_lines.append('## Body')
        md_lines.append('')

        if msg.htmlBody:
            # Convert HTML to plain text/markdown
            soup = BeautifulSoup(msg.htmlBody, 'html.parser')
            # Remove script and style elements
            for element in soup(['script', 'style']):
                element.decompose()
            body_text = soup.get_text(separator='\n').strip()
        else:
            body_text = msg.body or ''

        md_lines.append(body_text)
        md_lines.append('')

        # Attachments
        if msg.attachments:
            md_lines.append('## Attachments')
            md_lines.append('')
            for att in msg.attachments:
                size_kb = len(att.data) / 1024 if hasattr(att, 'data') and att.data else 0
                md_lines.append(f"- {att.longFilename or att.shortFilename or 'unnamed'} ({size_kb:.1f} KB)")
            md_lines.append('')

        msg.close()

        md_content = '\n'.join(md_lines)

        if output_path:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(md_content)
            print(f"Created: {output_path}")

        return md_content

    def msg_to_eml(self, input_path: str, output_path: str) -> None:
        """Convert Outlook MSG to standard EML format."""
        import extract_msg
        from email.mime.multipart import MIMEMultipart
        from email.mime.text import MIMEText
        from email.mime.base import MIMEBase
        from email import encoders
        from email.utils import formatdate

        msg = extract_msg.Message(input_path)

        # Create email message
        eml = MIMEMultipart()
        eml['Subject'] = msg.subject or ''
        eml['From'] = msg.sender or ''
        eml['To'] = msg.to or ''
        if msg.cc:
            eml['Cc'] = msg.cc
        if msg.date:
            eml['Date'] = str(msg.date)

        # Add body
        if msg.htmlBody:
            eml.attach(MIMEText(msg.htmlBody, 'html', 'utf-8'))
        elif msg.body:
            eml.attach(MIMEText(msg.body, 'plain', 'utf-8'))

        # Add attachments
        for att in msg.attachments:
            if hasattr(att, 'data') and att.data:
                part = MIMEBase('application', 'octet-stream')
                part.set_payload(att.data)
                encoders.encode_base64(part)
                filename = att.longFilename or att.shortFilename or 'attachment'
                part.add_header('Content-Disposition', f'attachment; filename="{filename}"')
                eml.attach(part)

        msg.close()

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(eml.as_string())

        print(f"Created: {output_path}")

    def msg_extract_attachments(self, input_path: str, output_dir: str) -> List[str]:
        """Extract all attachments from MSG file."""
        import extract_msg

        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)

        msg = extract_msg.Message(input_path)
        extracted = []

        for att in msg.attachments:
            filename = att.longFilename or att.shortFilename or 'attachment'
            file_path = output_path / filename

            # Handle duplicate filenames
            counter = 1
            while file_path.exists():
                stem = Path(filename).stem
                suffix = Path(filename).suffix
                file_path = output_path / f"{stem}_{counter}{suffix}"
                counter += 1

            if hasattr(att, 'data') and att.data:
                with open(file_path, 'wb') as f:
                    f.write(att.data)
                extracted.append(str(file_path))
                print(f"Extracted: {file_path}")

        msg.close()

        print(f"\nExtracted {len(extracted)} attachment(s) to {output_dir}")
        return extracted

    def _add_formatted_text(self, paragraph, text: str) -> None:
        """Add text with markdown formatting to a paragraph."""
        from docx.shared import Pt

        # Pattern to match formatting: **bold**, *italic*, `code`, ***bold italic***
        pattern = r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)'
        parts = re.split(pattern, text)

        for part in parts:
            if not part:
                continue

            if part.startswith('***') and part.endswith('***'):
                run = paragraph.add_run(part[3:-3])
                run.bold = True
                run.italic = True
            elif part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif part.startswith('`') and part.endswith('`'):
                run = paragraph.add_run(part[1:-1])
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
            else:
                paragraph.add_run(part)

    def _strip_formatting(self, text: str) -> str:
        """Remove markdown formatting from text."""
        text = re.sub(r'\*\*\*(.*?)\*\*\*', r'\1', text)
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'\*(.*?)\*', r'\1', text)
        text = re.sub(r'`(.*?)`', r'\1', text)
        return text

    def _extract_formatting(self, para) -> str:
        """Extract formatting from paragraph runs."""
        result = []
        for run in para.runs:
            text = run.text
            if run.bold and run.italic:
                text = f"***{text}***"
            elif run.bold:
                text = f"**{text}**"
            elif run.italic:
                text = f"*{text}*"
            result.append(text)
        return ''.join(result) or para.text

    def _table_to_markdown(self, table) -> str:
        """Convert DOCX table to markdown."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        if not rows:
            return ''

        # Build markdown table
        max_cols = max(len(r) for r in rows)
        lines = []

        # Header
        header = rows[0]
        header.extend([''] * (max_cols - len(header)))
        lines.append('| ' + ' | '.join(header) + ' |')

        # Separator
        lines.append('|' + '|'.join(['---'] * max_cols) + '|')

        # Data rows
        for row in rows[1:]:
            row.extend([''] * (max_cols - len(row)))
            lines.append('| ' + ' | '.join(row) + ' |')

        return '\n'.join(lines)

    def _extract_markdown_tables(self, md_content: str) -> list:
        """Extract all tables from markdown content."""
        tables = []
        lines = md_content.split('\n')
        current_table = []
        in_table = False

        for line in lines:
            # Check if line is a table row
            if '|' in line and line.strip().startswith('|'):
                cells = [c.strip() for c in line.strip().split('|')[1:-1]]
                # Skip separator rows (containing only -, :, and spaces)
                if all(set(c) <= set('-: ') for c in cells):
                    continue
                current_table.append(cells)
                in_table = True
            else:
                if in_table and current_table:
                    tables.append(current_table)
                    current_table = []
                in_table = False

        # Don't forget last table
        if current_table:
            tables.append(current_table)

        return tables


def main():
    """CLI entry point."""
    if len(sys.argv) < 2 or sys.argv[1] in ['-h', '--help']:
        print("Usage: convert.py <command> <input> [output] [options]")
        print("")
        print("Document Conversions:")
        print("  md2docx      Markdown → Word document")
        print("  docx2md      Word document → Markdown")
        print("  table2xlsx   Markdown tables → Excel spreadsheet")
        print("  xlsx2table   Excel spreadsheet → Markdown tables")
        print("  xlsx2csv     Excel spreadsheet → CSV")
        print("  csv2xlsx     CSV → Excel spreadsheet")
        print("")
        print("Excel Options (xlsx2table, xlsx2csv):")
        print("  --sheet NAME       Only convert specific sheet")
        print("  --sheets A,B,C     Convert multiple sheets (xlsx2table only)")
        print("  --range A1:D10     Extract specific cell range")
        print("  --range Sheet1!B2:F20  Range with sheet name")
        print("  --no-header        Don't treat first row as header")
        print("  --delimiter ';'    CSV delimiter (default: comma)")
        print("")
        print("Excel Info:")
        print("  xlsx-sheets  List all sheet names")
        print("  xlsx-info    Show sheet dimensions and ranges")
        print("")
        print("PDF Conversions:")
        print("  md2pdf       Markdown → PDF document")
        print("  pdf2md       PDF → Markdown (text extraction)")
        print("")
        print("PowerPoint Conversions:")
        print("  md2pptx      Markdown → PowerPoint presentation")
        print("  pptx2md      PowerPoint → Markdown")
        print("")
        print("Outlook MSG Conversions:")
        print("  msg2md       MSG email → Markdown")
        print("  msg2eml      MSG email → EML format")
        print("  msg-extract  Extract attachments from MSG")
        print("")
        print("Dependencies are auto-installed when needed.")
        print("")
        print("Examples:")
        print("  convert.py xlsx2table data.xlsx --sheet 'Sales Data'")
        print("  convert.py xlsx2table data.xlsx --range 'B2:F20'")
        print("  convert.py xlsx2table data.xlsx --range 'Q1!A1:D50' out.md")
        print("  convert.py xlsx-info report.xlsx")
        sys.exit(0 if len(sys.argv) >= 2 else 1)

    if len(sys.argv) < 3:
        print("Error: Missing input file")
        print("Usage: convert.py <command> <input> [output] [options]")
        sys.exit(1)

    command = sys.argv[1]
    input_path = sys.argv[2]

    # Parse remaining arguments
    remaining_args = sys.argv[3:]
    output_path = None
    sheets = None
    sheet = None
    range_notation = None
    no_header = False
    delimiter = ','

    i = 0
    while i < len(remaining_args):
        arg = remaining_args[i]
        if arg == '--sheet' and i + 1 < len(remaining_args):
            sheet = remaining_args[i + 1]
            sheets = [sheet]
            i += 2
        elif arg == '--sheets' and i + 1 < len(remaining_args):
            sheets = [s.strip() for s in remaining_args[i + 1].split(',')]
            i += 2
        elif arg == '--range' and i + 1 < len(remaining_args):
            range_notation = remaining_args[i + 1]
            i += 2
        elif arg == '--no-header':
            no_header = True
            i += 1
        elif arg == '--delimiter' and i + 1 < len(remaining_args):
            delimiter = remaining_args[i + 1]
            i += 2
        elif not arg.startswith('-') and output_path is None:
            output_path = arg
            i += 1
        else:
            i += 1

    # Auto-install dependencies
    ensure_dependencies(command)

    converter = DocumentConverter()

    # Word/Excel conversions
    if command == 'md2docx':
        if not output_path:
            output_path = Path(input_path).stem + '.docx'
        converter.md_to_docx(input_path, output_path)

    elif command == 'docx2md':
        if not output_path:
            output_path = Path(input_path).stem + '.md'
        converter.docx_to_md(input_path, output_path)

    elif command == 'table2xlsx':
        if not output_path:
            output_path = Path(input_path).stem + '.xlsx'
        converter.table_to_xlsx(input_path, output_path)

    elif command == 'xlsx2table':
        if not output_path:
            output_path = Path(input_path).stem + '_tables.md'
        converter.xlsx_to_markdown_table(
            input_path,
            output_path,
            sheets=sheets,
            range_notation=range_notation,
            no_header=no_header
        )

    elif command == 'xlsx-sheets':
        ensure_dependencies('xlsx2table')
        sheet_names = converter.list_sheets(input_path)
        print(f"Sheets in {input_path}:")
        for name in sheet_names:
            print(f"  - {name}")

    elif command == 'xlsx-info':
        ensure_dependencies('xlsx2table')
        info = converter.get_sheet_info(input_path)
        print(f"Sheet information for {input_path}:\n")
        for sheet_name, data in info.items():
            print(f"  {sheet_name}:")
            print(f"    Dimensions: {data['dimensions']}")
            print(f"    Rows: {data['max_row']}, Columns: {data['max_column']} ({data['max_column_letter']})")
            print()

    elif command == 'xlsx2csv':
        if not output_path:
            output_path = Path(input_path).stem + '.csv'
        converter.xlsx_to_csv(
            input_path,
            output_path,
            sheet=sheet,
            range_notation=range_notation,
            delimiter=delimiter
        )

    elif command == 'csv2xlsx':
        if not output_path:
            output_path = Path(input_path).stem + '.xlsx'
        converter.csv_to_xlsx(
            input_path,
            output_path,
            delimiter=delimiter,
            has_header=not no_header
        )

    # PDF conversions
    elif command == 'md2pdf':
        if not output_path:
            output_path = Path(input_path).stem + '.pdf'
        converter.md_to_pdf(input_path, output_path)

    elif command == 'pdf2md':
        if not output_path:
            output_path = Path(input_path).stem + '.md'
        converter.pdf_to_md(input_path, output_path)

    # PowerPoint conversions
    elif command == 'md2pptx':
        if not output_path:
            output_path = Path(input_path).stem + '.pptx'
        converter.md_to_pptx(input_path, output_path)

    elif command == 'pptx2md':
        if not output_path:
            output_path = Path(input_path).stem + '.md'
        converter.pptx_to_md(input_path, output_path)

    # MSG conversions
    elif command == 'msg2md':
        if not output_path:
            output_path = Path(input_path).stem + '.md'
        converter.msg_to_md(input_path, output_path)

    elif command == 'msg2eml':
        if not output_path:
            output_path = Path(input_path).stem + '.eml'
        converter.msg_to_eml(input_path, output_path)

    elif command == 'msg-extract':
        if not output_path:
            output_path = Path(input_path).stem + '_attachments'
        converter.msg_extract_attachments(input_path, output_path)

    else:
        print(f"Unknown command: {command}")
        print("Run 'convert.py --help' for available commands")
        sys.exit(1)


# Backwards compatibility alias
MarkdownConverter = DocumentConverter


if __name__ == '__main__':
    main()
